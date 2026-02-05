from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from math import ceil
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


INPUT_PATH = os.path.join("input", "presentation.pptx")
OUTPUT_PATH = os.path.join("output", "redesign.pptx")


# Slide geometry (16:9)
SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)

# Layout constants
MARGIN_X = Inches(0.75)
TITLE_Y = Inches(0.2)
TITLE_H = Inches(0.9)
CONTENT_Y = Inches(1.45)
FOOTER_H = Inches(0.45)
GAP = Inches(0.35)
STRIPE_W = Inches(0.12)
CARD_PAD_X = Inches(0.3)
CARD_PAD_Y = Inches(0.25)

# Colors
COLOR_BG_DARK = RGBColor(0x0A, 0x0A, 0x0A)
COLOR_TEXT_LIGHT = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_TEXT_DARK = RGBColor(0x0A, 0x0A, 0x0A)
COLOR_MUTED_LIGHT = RGBColor(0xC2, 0xC2, 0xC2)
COLOR_CARD = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_CARD_ALT = RGBColor(0xF1, 0xF5, 0xF9)
COLOR_ACCENT = RGBColor(0x5E, 0xC0, 0x62)
COLOR_ACCENT_2 = RGBColor(0x00, 0x97, 0xA7)

# Fonts
FONT_TITLE = "Montserrat"
FONT_BODY = "Inter"


@dataclass
class SlideData:
    title: str
    lead: list[str]
    bullets: list[str]
    callouts: list[str]
    cta: list[str]
    contacts: list[str]
    images: list[dict]


def _get_max_font_size(shape) -> float:
    max_size = 0.0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                max_size = max(max_size, run.font.size.pt)
    return max_size


def _classify_contact(text: str) -> str | None:
    if "Москва" in text:
        return "address"
    if text.strip().startswith("@"):
        return "handle"
    if "+7" in text:
        return "phone"
    if "@" in text and "." in text:
        return "email"
    if ".ru" in text or "asg-" in text:
        return "site"
    return None


def _is_callout(text: str) -> bool:
    stripped = text.strip()
    if stripped.startswith("Вывод:"):
        return True
    if stripped.startswith("ЭТО ") or stripped.startswith("ВЫ "):
        return True
    letters = [c for c in stripped if c.isalpha()]
    if not letters:
        return False
    upper = [c for c in letters if c.upper() == c]
    if len(upper) / len(letters) > 0.85 and len(stripped) > 20:
        return True
    if "→" in stripped and len(stripped) > 10:
        return True
    return False


def _is_cta(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if len(stripped) > 30:
        return False
    if stripped.upper() != stripped:
        return False
    return "ЗАЯВ" in stripped or "ВСТРЕЧ" in stripped


def _split_items(text: str) -> list[str]:
    items: list[str] = []
    for chunk in text.split("|"):
        for part in chunk.split("\n"):
            item = part.strip()
            if item:
                items.append(item)
    return items


def _extract_slide_data(slide) -> SlideData:
    texts = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if not text:
            continue
        texts.append(
            {
                "text": text,
                "top": shape.top,
                "left": shape.left,
                "size": _get_max_font_size(shape),
            }
        )
    texts.sort(key=lambda t: (t["top"], t["left"]))

    images = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images.append(
                {
                    "blob": shape.image.blob,
                    "width": int(shape.width),
                    "height": int(shape.height),
                    "area": int(shape.width) * int(shape.height),
                }
            )

    title = ""
    if texts:
        texts_by_size = sorted(texts, key=lambda t: (-t["size"], t["top"], t["left"]))
        title = texts_by_size[0]["text"]

    contacts = []
    content_texts = []
    for item in texts:
        if item["text"] == title:
            continue
        if _classify_contact(item["text"]):
            contacts.append(item["text"])
        else:
            content_texts.append(item["text"])

    lead: list[str] = []
    bullets: list[str] = []
    callouts: list[str] = []
    cta: list[str] = []

    for text in content_texts:
        if _is_cta(text):
            cta.append(text)
            continue
        if _is_callout(text):
            callouts.append(text)
            continue
        if text.strip().endswith(":"):
            lead.append(text)
            continue
        if text.strip().endswith(".") and len(text) > 40 and not bullets:
            lead.append(text)
            continue
        bullets.extend(_split_items(text))

    return SlideData(
        title=title,
        lead=lead,
        bullets=bullets,
        callouts=callouts,
        cta=cta,
        contacts=contacts,
        images=images,
    )


def _add_rect(slide, left, top, width, height, fill_color, line_color=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def _add_text_box(
    slide,
    left,
    top,
    width,
    height,
    lines,
    font_name,
    font_size,
    bold=False,
    color=COLOR_TEXT_DARK,
    align=PP_ALIGN.LEFT,
    line_spacing=1.1,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font_name
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
    return box


def _apply_text_frame(tf, lines, font_name, font_size, bold=False, color=COLOR_TEXT_DARK, align=PP_ALIGN.LEFT):
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font_name
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.alignment = align


def _add_picture_contain(slide, image_blob, box_left, box_top, box_width, box_height):
    stream = BytesIO(image_blob)
    image = Image.open(stream)
    img_w, img_h = image.size
    stream.seek(0)
    if img_w == 0 or img_h == 0:
        return
    scale = min(box_width / img_w, box_height / img_h)
    width = int(img_w * scale)
    height = int(img_h * scale)
    left = int(box_left + (box_width - width) / 2)
    top = int(box_top + (box_height - height) / 2)
    slide.shapes.add_picture(stream, left, top, width=width, height=height)


def _pick_accents(index: int) -> tuple[RGBColor, RGBColor]:
    if index % 2 == 0:
        return COLOR_ACCENT, COLOR_ACCENT_2
    return COLOR_ACCENT_2, COLOR_ACCENT


def _add_base(slide, accent, accent_2):
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, COLOR_BG_DARK)
    _add_rect(slide, 0, 0, Inches(0.28), SLIDE_H, accent)
    _add_rect(slide, Inches(0.28), 0, SLIDE_W, Inches(0.08), accent_2)
    diagonal = _add_rect(slide, Inches(6.6), Inches(-0.7), Inches(4.6), Inches(1.9), accent_2)
    diagonal.rotation = -8


def _add_title(slide, title: str, accent):
    _add_text_box(
        slide,
        MARGIN_X,
        TITLE_Y,
        SLIDE_W - 2 * MARGIN_X,
        TITLE_H,
        [title],
        FONT_TITLE,
        30,
        bold=True,
        color=COLOR_TEXT_LIGHT,
        line_spacing=1.0,
    )
    _add_rect(slide, MARGIN_X, TITLE_Y + TITLE_H - Inches(0.08), Inches(2.6), Inches(0.08), accent)


def _format_contacts(contacts: list[str]) -> str:
    ordered = {"address": [], "handle": [], "phone": [], "email": [], "site": []}
    for text in contacts:
        kind = _classify_contact(text)
        if not kind:
            continue
        ordered[kind].append(text)
    combined = []
    for key in ["address", "handle", "phone", "email", "site"]:
        combined.extend(ordered[key])
    if not combined:
        return ""
    return " | ".join(combined)


def _add_footer(slide, contacts: list[str], accent_2):
    text = _format_contacts(contacts)
    if not text:
        return
    _add_rect(slide, 0, SLIDE_H - Inches(0.04), SLIDE_W, Inches(0.04), accent_2)
    _add_text_box(
        slide,
        MARGIN_X,
        SLIDE_H - FOOTER_H + Inches(0.06),
        SLIDE_W - 2 * MARGIN_X,
        FOOTER_H - Inches(0.12),
        [text],
        FONT_BODY,
        10,
        color=COLOR_MUTED_LIGHT,
    )


def _add_card(slide, left, top, width, height, accent):
    card = _add_rect(slide, left, top, width, height, COLOR_CARD, radius=True)
    _add_rect(slide, left, top, STRIPE_W, height, accent)
    return card


def _card_inner_bounds(left, top, width, height):
    inner_left = left + STRIPE_W + CARD_PAD_X
    inner_top = top + CARD_PAD_Y
    inner_width = width - STRIPE_W - (CARD_PAD_X * 2)
    inner_height = height - (CARD_PAD_Y * 2)
    return inner_left, inner_top, inner_width, inner_height


def _add_callout(slide, left, top, width, height, lines, accent):
    rect = _add_rect(slide, left, top, width, height, accent, line_color=accent, radius=True)
    _apply_text_frame(
        rect.text_frame,
        lines,
        FONT_TITLE,
        14,
        bold=True,
        color=COLOR_TEXT_LIGHT,
        align=PP_ALIGN.CENTER,
    )


def _largest_image(images: list[dict], min_ratio: float = 0.1) -> dict | None:
    if not images:
        return None
    slide_area = int(SLIDE_W) * int(SLIDE_H)
    images_sorted = sorted(images, key=lambda img: img["area"], reverse=True)
    if images_sorted[0]["area"] >= slide_area * min_ratio:
        return images_sorted[0]
    return None


def _build_cover_slide(slide, data: SlideData, accent, accent_2):
    _add_base(slide, accent, accent_2)
    title_text = data.title.replace(" | ", "\n")
    _add_text_box(
        slide,
        MARGIN_X,
        Inches(0.55),
        Inches(5.4),
        Inches(2.1),
        [title_text],
        FONT_TITLE,
        36,
        bold=True,
        color=COLOR_TEXT_LIGHT,
        line_spacing=1.05,
    )

    subtitle_source = data.lead if data.lead else data.callouts
    subtitle = [t.strip("-") for t in subtitle_source] if subtitle_source else []
    if subtitle:
        _add_text_box(
            slide,
            MARGIN_X,
            Inches(2.45),
            Inches(5.0),
            Inches(0.9),
            subtitle,
            FONT_BODY,
            16,
            color=COLOR_MUTED_LIGHT,
        )

    logo = _largest_image(data.images, min_ratio=0.02)
    card_left = Inches(6.2)
    card_top = Inches(1.3)
    card_w = Inches(3.1)
    card_h = Inches(2.5)
    _add_card(slide, card_left, card_top, card_w, card_h, accent_2)
    if logo:
        _add_picture_contain(
            slide,
            logo["blob"],
            card_left + STRIPE_W + Inches(0.2),
            card_top + Inches(0.2),
            card_w - STRIPE_W - Inches(0.4),
            card_h - Inches(0.4),
        )

    _add_footer(slide, data.contacts, accent_2)


def _build_text_slide(slide, data: SlideData, accent, accent_2):
    _add_base(slide, accent, accent_2)
    _add_title(slide, data.title, accent)
    card_left = MARGIN_X
    card_top = CONTENT_Y
    card_w = SLIDE_W - 2 * MARGIN_X
    card_h = SLIDE_H - CONTENT_Y - FOOTER_H - Inches(0.12)
    _add_card(slide, card_left, card_top, card_w, card_h, accent)

    inner_left, inner_top, inner_w, inner_h = _card_inner_bounds(card_left, card_top, card_w, card_h)
    cursor_y = inner_top

    if data.lead:
        lead_h = Inches(0.28 * len(data.lead) + 0.15)
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            lead_h,
            data.lead,
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )
        cursor_y += lead_h + Inches(0.1)

    callout_h = Inches(0.6) if data.callouts else Inches(0)
    bullets_h = inner_h - (cursor_y - inner_top) - callout_h - Inches(0.05)
    bullets = data.bullets

    if bullets:
        if len(bullets) > 6:
            mid = ceil(len(bullets) / 2)
            left_items = [f"• {b}" for b in bullets[:mid]]
            right_items = [f"• {b}" for b in bullets[mid:]]
            left_w = (inner_w - GAP) / 2
            right_x = inner_left + left_w + GAP
            _add_text_box(
                slide,
                inner_left,
                cursor_y,
                left_w,
                bullets_h,
                left_items,
                FONT_BODY,
                16,
                color=COLOR_TEXT_DARK,
            )
            _add_text_box(
                slide,
                right_x,
                cursor_y,
                left_w,
                bullets_h,
                right_items,
                FONT_BODY,
                16,
                color=COLOR_TEXT_DARK,
            )
        else:
            _add_text_box(
                slide,
                inner_left,
                cursor_y,
                inner_w,
                bullets_h,
                [f"• {b}" for b in bullets],
                FONT_BODY,
                16,
                color=COLOR_TEXT_DARK,
            )

    if data.callouts:
        callout_top = inner_top + inner_h - callout_h
        _add_callout(slide, inner_left, callout_top, inner_w, callout_h, data.callouts, accent_2)

    _add_footer(slide, data.contacts, accent_2)


def _build_split_slide(slide, data: SlideData, image: dict | None, accent, accent_2):
    _add_base(slide, accent, accent_2)
    _add_title(slide, data.title, accent)
    card_top = CONTENT_Y
    card_h = SLIDE_H - CONTENT_Y - FOOTER_H - Inches(0.12)
    card_w = SLIDE_W - 2 * MARGIN_X
    left_w = (card_w - GAP) / 2
    right_w = left_w
    left_x = MARGIN_X
    right_x = left_x + left_w + GAP

    _add_card(slide, left_x, card_top, left_w, card_h, accent)
    _add_card(slide, right_x, card_top, right_w, card_h, accent_2)

    inner_left, inner_top, inner_w, inner_h = _card_inner_bounds(left_x, card_top, left_w, card_h)
    cursor_y = inner_top

    if data.lead:
        lead_h = Inches(0.28 * len(data.lead) + 0.15)
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            lead_h,
            data.lead,
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )
        cursor_y += lead_h + Inches(0.1)

    if data.bullets:
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            inner_h - (cursor_y - inner_top),
            [f"• {b}" for b in data.bullets],
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )

    if data.callouts:
        _add_callout(
            slide,
            inner_left,
            inner_top + inner_h - Inches(0.6),
            inner_w,
            Inches(0.6),
            data.callouts,
            accent,
        )

    if image:
        img_left, img_top, img_w, img_h = _card_inner_bounds(right_x, card_top, right_w, card_h)
        _add_picture_contain(slide, image["blob"], img_left, img_top, img_w, img_h)

    _add_footer(slide, data.contacts, accent_2)


def _build_image_slide(slide, data: SlideData, image: dict | None, accent, accent_2):
    _add_base(slide, accent, accent_2)
    _add_title(slide, data.title, accent)
    card_left = MARGIN_X
    card_top = CONTENT_Y
    card_w = SLIDE_W - 2 * MARGIN_X
    card_h = SLIDE_H - CONTENT_Y - FOOTER_H - Inches(0.12)
    _add_card(slide, card_left, card_top, card_w, card_h, accent)

    inner_left, inner_top, inner_w, inner_h = _card_inner_bounds(card_left, card_top, card_w, card_h)
    cursor_y = inner_top

    if data.lead:
        lead_h = Inches(0.28 * len(data.lead) + 0.15)
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            lead_h,
            data.lead,
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )
        cursor_y += lead_h + Inches(0.1)

    if data.callouts:
        _add_callout(slide, inner_left, cursor_y, inner_w, Inches(0.6), data.callouts, accent_2)
        cursor_y += Inches(0.7)

    if image:
        _add_picture_contain(
            slide,
            image["blob"],
            inner_left,
            cursor_y,
            inner_w,
            inner_h - (cursor_y - inner_top),
        )

    _add_footer(slide, data.contacts, accent_2)


def _build_cta_slide(slide, data: SlideData, accent, accent_2):
    _add_base(slide, accent, accent_2)
    _add_title(slide, data.title, accent)
    card_left = MARGIN_X
    card_top = CONTENT_Y
    card_w = SLIDE_W - 2 * MARGIN_X
    card_h = SLIDE_H - CONTENT_Y - FOOTER_H - Inches(0.12)
    _add_card(slide, card_left, card_top, card_w, card_h, accent)

    inner_left, inner_top, inner_w, inner_h = _card_inner_bounds(card_left, card_top, card_w, card_h)
    cursor_y = inner_top

    if data.lead:
        lead_h = Inches(0.28 * len(data.lead) + 0.15)
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            lead_h,
            data.lead,
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )
        cursor_y += lead_h + Inches(0.1)

    if data.bullets:
        _add_text_box(
            slide,
            inner_left,
            cursor_y,
            inner_w,
            inner_h - Inches(1.1),
            [f"• {b}" for b in data.bullets],
            FONT_BODY,
            16,
            color=COLOR_TEXT_DARK,
        )

    if data.cta:
        cta_w = Inches(3.0)
        cta_h = Inches(0.62)
        start_x = inner_left
        cta_top = inner_top + inner_h - Inches(0.8)
        for idx, text in enumerate(data.cta[:2]):
            color = accent if idx == 0 else accent_2
            rect = _add_rect(
                slide,
                start_x + idx * (cta_w + Inches(0.4)),
                cta_top,
                cta_w,
                cta_h,
                color,
                line_color=color,
                radius=True,
            )
            _apply_text_frame(
                rect.text_frame,
                [text],
                FONT_TITLE,
                14,
                bold=True,
                color=COLOR_TEXT_LIGHT,
                align=PP_ALIGN.CENTER,
            )

    _add_footer(slide, data.contacts, accent_2)


def main():
    prs = Presentation(INPUT_PATH)
    new_prs = Presentation()
    new_prs.slide_width = SLIDE_W
    new_prs.slide_height = SLIDE_H
    blank = new_prs.slide_layouts[6]

    for index, slide in enumerate(prs.slides):
        data = _extract_slide_data(slide)
        new_slide = new_prs.slides.add_slide(blank)
        accent, accent_2 = _pick_accents(index)

        if index == 0:
            _build_cover_slide(new_slide, data, accent, accent_2)
            continue

        if data.cta:
            _build_cta_slide(new_slide, data, accent, accent_2)
            continue

        large_image = _largest_image(data.images, min_ratio=0.12)

        if large_image and len(data.bullets) <= 2 and len(data.lead) <= 1:
            _build_image_slide(new_slide, data, large_image, accent, accent_2)
        elif large_image:
            _build_split_slide(new_slide, data, large_image, accent, accent_2)
        else:
            _build_text_slide(new_slide, data, accent, accent_2)

    new_prs.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
